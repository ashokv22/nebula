import csv
import io
import zipfile

from openpyxl import load_workbook

from cargo.schemas import (
    MAX_CELL_CHARS,
    MAX_COLUMNS,
    MAX_CONTENT_LENGTH,
    MAX_ROWS,
    MAX_ZIP_FILES,
    MAX_ZIP_RATIO,
    MAX_ZIP_UNCOMPRESSED,
    FileProblem,
)


def check_xlsx_archive(blob):
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            files = z.infolist()
            if len(files) > MAX_ZIP_FILES:
                raise FileProblem('Workbook contains too many archive entries')
            if sum(f.file_size for f in files) > MAX_ZIP_UNCOMPRESSED:
                raise FileProblem('Workbook expands beyond safe size')
            for f in files:
                if f.flag_bits & 1:
                    raise FileProblem('Encrypted workbooks are unsupported')
                if f.file_size > MAX_ZIP_UNCOMPRESSED or (f.file_size > 1024 * 1024 and f.file_size > max(f.compress_size, 1) * MAX_ZIP_RATIO):
                    raise FileProblem('Workbook contains an oversized compressed entry')
                name = f.filename.replace('\\', '/')
                if name.startswith('/') or '..' in name.split('/'):
                    raise FileProblem('Unsafe workbook archive path')
                if name.lower().endswith(('vbaproject.bin', '.bin')):
                    raise FileProblem('Workbook binary components are unsupported')
            names = {f.filename for f in files}
            if '[Content_Types].xml' not in names or 'xl/workbook.xml' not in names:
                raise FileProblem('Not a valid XLSX workbook')
    except (zipfile.BadZipFile, OSError, ValueError) as exc:
        raise FileProblem('Invalid XLSX archive') from exc


def parse_xlsx(blob):
    check_xlsx_archive(blob)
    try:
        wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=False, keep_links=False)
        try:
            if len(wb.sheetnames) != 1:
                raise FileProblem('Exactly one worksheet is required')
            ws = wb.active
            if ws.max_row and ws.max_row > MAX_ROWS + 1:
                raise FileProblem('Maximum 1000 data rows allowed')
            if ws.max_column and ws.max_column > MAX_COLUMNS:
                raise FileProblem('Too many columns')
            rows = []
            for row in ws.iter_rows(max_row=MAX_ROWS + 2, max_col=MAX_COLUMNS + 1):
                values = []
                for cell in row:
                    if cell.data_type == 'f':
                        raise FileProblem('Formula cells are not supported')
                    value = '' if cell.value is None else str(cell.value).strip()
                    if len(value) > MAX_CELL_CHARS:
                        raise FileProblem('Cell value exceeds 500 characters')
                    values.append(value)
                while values and not values[-1]:
                    values.pop()
                rows.append(values)
            if len(rows) > MAX_ROWS + 1:
                raise FileProblem('Maximum 1000 data rows allowed')
            return rows
        finally:
            wb.close()
    except FileProblem:
        raise
    except Exception as exc:
        raise FileProblem('Could not read XLSX workbook') from exc


def parse_csv(blob):
    try:
        content = blob.decode('utf-8-sig', errors='strict')
    except UnicodeError as exc:
        raise FileProblem('CSV must be UTF-8 encoded') from exc
    try:
        rows = []
        for row in csv.reader(io.StringIO(content, newline=''), strict=True):
            if len(row) > MAX_COLUMNS:
                raise FileProblem('Too many columns')
            clean = [value.strip() for value in row]
            if any(len(v) > MAX_CELL_CHARS for v in clean):
                raise FileProblem('Cell value exceeds 500 characters')
            rows.append(clean)
            if len(rows) > MAX_ROWS + 1:
                raise FileProblem('Maximum 1000 data rows allowed')
        return rows
    except csv.Error as exc:
        raise FileProblem('Malformed CSV') from exc


def read_uploaded_rows(filename, blob):
    if filename is None:
        raise FileProblem('Multipart form field "file" is required')
    filename = filename.lower()
    if not filename.endswith(('.xlsx', '.csv')):
        raise FileProblem('Only .xlsx and .csv are supported')
    if blob is None:
        raise FileProblem('File is empty')
    if len(blob) > MAX_CONTENT_LENGTH:
        raise FileProblem('Upload exceeds 2 MB', 413)
    if not blob:
        raise FileProblem('File is empty')
    if filename.endswith('.xlsx'):
        return parse_xlsx(blob)
    if blob.startswith((b'PK\x03\x04', b'MZ')):
        raise FileProblem('File content does not match CSV format')
    return parse_csv(blob)